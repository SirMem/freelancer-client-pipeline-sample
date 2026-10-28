#!/usr/bin/env python3
"""Build Freelancer-Client-Pipeline-Sample.xlsx — free structural seam of client-tracker-v1."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from datetime import date
from pathlib import Path

OUT = Path(__file__).resolve().parent / "Freelancer-Client-Pipeline-Sample.xlsx"

H = Font(bold=True, color="FFFFFF")
HF = PatternFill("solid", fgColor="2F4F4F")
IN = PatternFill("solid", fgColor="FFF8DC")
T = Font(bold=True, size=16)
B = Font(bold=True)


def head(ws, row, cols, widths):
    for i, (c, w) in enumerate(zip(cols, widths), 1):
        x = ws.cell(row, i, c)
        x.font = H
        x.fill = HF
        x.alignment = Alignment(horizontal="center", wrap_text=True)
        ws.column_dimensions[x.column_letter].width = w
    ws.freeze_panes = ws.cell(row + 1, 1)


def build():
    wb = Workbook()
    s = wb.active
    s.title = "Start Here"
    lines = [
        "Freelancer Client & Project Pipeline — Free Sample",
        "",
        "What this is",
        "A complete, usable mini CRM pipeline: client roster + projects with status, fee type, auto Value, Paid?, due date, and Overdue flag. Same quality as the paid tracker’s Clients + Projects core — not a crippled demo.",
        "",
        "What you can do in 10 minutes",
        "1) Skim Clients — replace EXAMPLE rows with your own roster (Active? Yes/No).",
        "2) On Projects, pick a client, set Status (Lead / Active / Waiting / Done / Lost) and Fee type (Hourly or Flat).",
        "3) Fill Rate + Est. hours (Hourly) or Flat fee (Flat). Value (auto) computes. Mark Paid? when money arrives; Overdue? flags past-due open work.",
        "4) Read the Pipeline snapshot at the bottom of Projects — Active count, pipeline $, won $.",
        "",
        "Seam map (what this deliberately does NOT do)",
        "FREE sample includes: Start Here + Clients (small roster) + Projects pipeline (Status / Fee type / Value / Paid? / Due / Overdue) + snapshot totals.",
        "NOT in free sample (paid Freelancer Client & Project Tracker breadth):",
        "  • Settings sheet (currency symbol, default hourly rate, studio name)",
        "  • Time Log (hours roll into Hours logged on each project)",
        "  • Dashboard (active / overdue / pipeline / won KPIs + by-status + by-client rollups)",
        "  • Default-rate fallback when a project Rate cell is blank",
        "",
        "Gap questions (reply with one word — this is the real demand signal)",
        "Downloads prove the funnel exists. A specific question proves purchase intent. Pick one and reply (email / DM / GitHub issue):",
        "1) PIPELINE — Do you want a Dashboard of pipeline vs won (by status + by client) without building a second sheet?",
        "2) TIMELOG — Do you want hours logged against a project to roll into Hours logged automatically?",
        "3) SETTINGS — Do you want default rate / currency / studio settings so blank project rates inherit a studio default?",
        "If none apply, keep the free sheet — no pitch. If one word fits, that is the paid Freelancer Client & Project Tracker seam.",
        "",
        "What “missing” looks like (not a vague “buy full pack”)",
        "Free sample solves: track clients and a status pipeline with fee math. Paid adds: Settings + Time Log + Dashboard so studio KPIs and hour rollups live in one file.",
        "",
        "License",
        "Free for personal use. Example data is fictional. Method: standard freelance project-tracking practice (client roster + status pipeline + fee type + due dates). Not accounting software or tax advice. Gap-question CTA added 2026-10-28 (free-client-sample-v1).",
    ]
    for i, l in enumerate(lines, 1):
        s.cell(i, 1, l)
    s["A1"].font = T
    s.column_dimensions["A"].width = 120

    # Clients — small roster
    cl = wb.create_sheet("Clients")
    head(cl, 1, ["Client", "Email", "Company", "Notes", "Active?"], [26, 30, 24, 36, 10])
    dva = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True)
    cl.add_data_validation(dva)
    clients = [
        ("EXAMPLE Acme Co", "ap@acme.example", "Acme Co", "retainer prospect", "Yes"),
        ("EXAMPLE Blue Fox Media", "hello@bluefox.example", "Blue Fox", "pays on Net 14", "Yes"),
        ("EXAMPLE Nora Lee", "nora@example.com", "", "one-off brand work", "Yes"),
        ("EXAMPLE Quiet Lab", "ops@quietlab.example", "Quiet Lab", "paused", "No"),
    ]
    for r in range(2, 52):
        for c in range(1, 6):
            cl.cell(r, c).fill = IN
        dva.add(f"E{r}")
    for r, row in enumerate(clients, 2):
        for c, v in enumerate(row, 1):
            cl.cell(r, c, v)

    # Projects — pipeline (no Hours logged / no Settings dependency)
    pr = wb.create_sheet("Projects")
    head(
        pr,
        1,
        [
            "Project",
            "Client",
            "Status",
            "Fee type",
            "Rate",
            "Est. hours",
            "Flat fee",
            "Value (auto)",
            "Start",
            "Due date",
            "Paid?",
            "Notes",
            "Overdue?",
        ],
        [28, 24, 12, 10, 10, 10, 12, 12, 12, 12, 10, 28, 10],
    )
    dvc = DataValidation(type="list", formula1="=Clients!$A$2:$A$50", allow_blank=True)
    dvs = DataValidation(type="list", formula1='"Lead,Active,Waiting,Done,Lost"', allow_blank=True)
    dvf = DataValidation(type="list", formula1='"Hourly,Flat"', allow_blank=True)
    dvp = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True)
    for dv in (dvc, dvs, dvf, dvp):
        pr.add_data_validation(dv)

    # Project, Client, Status, Fee type, Rate, Est hours, Flat fee, Start, Due, Paid?, Notes
    samples = [
        (
            "EXAMPLE Brand refresh",
            "EXAMPLE Acme Co",
            "Active",
            "Flat",
            None,
            None,
            2400,
            date(2026, 9, 15),
            date(2026, 11, 15),
            "No",
            "kickoff done",
        ),
        (
            "EXAMPLE Monthly retainer",
            "EXAMPLE Blue Fox Media",
            "Active",
            "Hourly",
            85,
            20,
            None,
            date(2026, 10, 1),
            date(2026, 10, 31),
            "No",
            "Oct hours",
        ),
        (
            "EXAMPLE Landing page copy",
            "EXAMPLE Nora Lee",
            "Waiting",
            "Flat",
            None,
            None,
            650,
            date(2026, 10, 5),
            date(2026, 10, 20),
            "No",
            "awaiting feedback",
        ),
        (
            "EXAMPLE Logo concepts",
            "EXAMPLE Quiet Lab",
            "Lead",
            "Flat",
            None,
            None,
            900,
            date(2026, 10, 22),
            date(2026, 11, 30),
            "No",
            "proposal sent",
        ),
        (
            "EXAMPLE Social kit",
            "EXAMPLE Acme Co",
            "Done",
            "Hourly",
            75,
            12,
            None,
            date(2026, 8, 1),
            date(2026, 8, 28),
            "Yes",
            "invoiced & paid",
        ),
        (
            "EXAMPLE Pitch deck",
            "EXAMPLE Nora Lee",
            "Lost",
            "Flat",
            None,
            None,
            1200,
            date(2026, 7, 10),
            date(2026, 8, 1),
            "No",
            "chose another vendor",
        ),
    ]
    n = len(samples)

    for r in range(2, 52):
        for c in (1, 2, 3, 4, 5, 6, 7, 9, 10, 11, 12):
            pr.cell(r, c).fill = IN
        dvc.add(f"B{r}")
        dvs.add(f"C{r}")
        dvf.add(f"D{r}")
        dvp.add(f"K{r}")
        # Value: Hourly = Rate*Est hours (no Settings fallback — Rate required); Flat = Flat fee
        pr.cell(
            r,
            8,
            f'=IF(A{r}="","",IF(D{r}="Hourly",IF(E{r}="",0,E{r})*IF(F{r}="",0,F{r}),IF(D{r}="Flat",IF(G{r}="",0,G{r}),0)))',
        ).number_format = "#,##0.00"
        # Overdue?: Due date passed AND status in Lead/Active/Waiting
        pr.cell(
            r,
            13,
            f'=IF(OR(A{r}="",J{r}=""),"",IF(AND(TODAY()>J{r},OR(C{r}="Lead",C{r}="Active",C{r}="Waiting")),"Yes","No"))',
        )
        for c in (5, 7):
            pr.cell(r, c).number_format = "#,##0.00"
        pr.cell(r, 6).number_format = "0.00"
        for c in (9, 10):
            pr.cell(r, c).number_format = "yyyy-mm-dd"

    for r, row in enumerate(samples, 2):
        pr.cell(r, 1, row[0])
        pr.cell(r, 2, row[1])
        pr.cell(r, 3, row[2])
        pr.cell(r, 4, row[3])
        if row[4] is not None:
            pr.cell(r, 5, row[4])
        if row[5] is not None:
            pr.cell(r, 6, row[5])
        if row[6] is not None:
            pr.cell(r, 7, row[6])
        pr.cell(r, 9, row[7])
        pr.cell(r, 10, row[8])
        pr.cell(r, 11, row[9])
        pr.cell(r, 12, row[10])

    red = PatternFill("solid", fgColor="F8D7DA")
    grn = PatternFill("solid", fgColor="D4EDDA")
    ylw = PatternFill("solid", fgColor="FFF3CD")
    pr.conditional_formatting.add("C2:C51", FormulaRule(formula=['$C2="Done"'], fill=grn))
    pr.conditional_formatting.add("C2:C51", FormulaRule(formula=['$C2="Lost"'], fill=red))
    pr.conditional_formatting.add("C2:C51", FormulaRule(formula=['$C2="Waiting"'], fill=ylw))
    pr.conditional_formatting.add("M2:M51", FormulaRule(formula=['$M2="Yes"'], fill=red))

    # Pipeline snapshot (auto) — mini totals, not a full Dashboard
    last = 1 + n
    pr.cell(54, 1, "Pipeline snapshot (auto)").font = B
    pr.cell(55, 1, "Active projects (count)")
    pr.cell(55, 2, f'=COUNTIF(C2:C{last},"Active")').number_format = "0"
    pr.cell(56, 1, "Waiting (count)")
    pr.cell(56, 2, f'=COUNTIF(C2:C{last},"Waiting")').number_format = "0"
    pr.cell(57, 1, "Leads (count)")
    pr.cell(57, 2, f'=COUNTIF(C2:C{last},"Lead")').number_format = "0"
    pr.cell(58, 1, "Overdue (count)")
    pr.cell(58, 2, f'=COUNTIF(M2:M{last},"Yes")').number_format = "0"
    pr.cell(59, 1, "Pipeline value (Lead+Active+Waiting)")
    pr.cell(
        59,
        2,
        f'=SUMIF(C2:C{last},"Lead",H2:H{last})+SUMIF(C2:C{last},"Active",H2:H{last})+SUMIF(C2:C{last},"Waiting",H2:H{last})',
    ).number_format = "#,##0.00"
    pr.cell(60, 1, "Won revenue (Done)")
    pr.cell(60, 2, f'=SUMIF(C2:C{last},"Done",H2:H{last})').number_format = "#,##0.00"
    pr.cell(61, 1, "Paid (Done + Paid Yes)")
    pr.cell(61, 2, f'=SUMIFS(H2:H{last},C2:C{last},"Done",K2:K{last},"Yes")').number_format = "#,##0.00"
    pr.cell(62, 1, "Lost pipeline (Lost)")
    pr.cell(62, 2, f'=SUMIF(C2:C{last},"Lost",H2:H{last})').number_format = "#,##0.00"

    pr.cell(64, 1, "Tip: delete EXAMPLE rows after you paste your own. Yellow = inputs. Value / Overdue are formulas.")
    pr.cell(
        65,
        1,
        "This sample is Clients + Projects pipeline only. Paid Freelancer Client & Project Tracker adds Settings + Time Log + Dashboard.",
    )

    wb.save(OUT)
    print("ok", OUT)
    print("hand-calc expected: Active=2 Waiting=1 Lead=1 pipeline=5650 won=900 paid=900 lost=1200")
    print("overdue: Landing page (due 2026-10-20, Waiting) = Yes if TODAY()>2026-10-20")


if __name__ == "__main__":
    build()
